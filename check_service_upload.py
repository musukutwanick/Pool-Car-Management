"""
Diagnostic script to check what's in the Excel file for service upload.
Usage: python check_service_upload.py <path_to_excel_file>
"""

import sys
import openpyxl

if len(sys.argv) < 2:
    print("Usage: python check_service_upload.py <path_to_excel_file>")
    sys.exit(1)

file_path = sys.argv[1]

print(f"\n=== Analyzing Excel File: {file_path} ===\n")

try:
    # Load workbook
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    ws = wb.active
    
    print(f"Worksheet Name: {ws.title}")
    print(f"Total Rows: {ws.max_row}")
    print(f"Total Columns: {ws.max_column}\n")
    
    # Show header row
    print("=== HEADER ROW (Row 1) ===")
    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    for col_idx, header in enumerate(header_row, start=1):
        print(f"Column {col_idx} ({chr(64+col_idx)}): {header}")
    
    print("\n=== DATA ROWS (First 10 rows) ===")
    # Show first 10 data rows
    for row_num, row in enumerate(ws.iter_rows(min_row=2, max_row=11, values_only=True), start=2):
        print(f"\n--- Row {row_num} ---")
        if all(cell is None or str(cell).strip() == '' for cell in row):
            print("  [EMPTY ROW - would be skipped]")
            continue
        
        for col_idx, cell_value in enumerate(row, start=1):
            if cell_value is not None and str(cell_value).strip() != '':
                print(f"  Column {chr(64+col_idx)}: {cell_value}")
    
    print("\n=== SUMMARY ===")
    non_empty_rows = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not all(cell is None or str(cell).strip() == '' for cell in row):
            non_empty_rows += 1
    
    print(f"Total non-empty data rows: {non_empty_rows}")
    
    # Check for registration numbers
    print("\n=== CHECKING REGISTRATION NUMBERS (Column B) ===")
    reg_numbers = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell is None or str(cell).strip() == '' for cell in row):
            continue
        reg = row[1] if len(row) > 1 else None
        if reg and str(reg).strip():
            reg_clean = str(reg).strip()
            if reg_clean.lower() not in ('none', 'n/a', '', 'registration number'):
                reg_numbers.append((row_num, reg_clean))
                print(f"  Row {row_num}: {reg_clean}")
            else:
                print(f"  Row {row_num}: INVALID - '{reg_clean}' (would be skipped)")
    
    print(f"\nTotal valid registration numbers found: {len(reg_numbers)}")
    
except Exception as e:
    print(f"ERROR: {str(e)}")
    import traceback
    traceback.print_exc()

print("\n=== END OF ANALYSIS ===\n")
